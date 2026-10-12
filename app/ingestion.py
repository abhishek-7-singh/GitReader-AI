import os
import subprocess
from pathlib import Path
from typing import Dict, Any, List, Optional
from datetime import datetime, timezone

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader
from docx import Document as WordDocument

from .config import MAX_FILE_BYTES
from .db import upsert_repository, record_commit, record_change_log, record_document
from .rag import add_documents, delete_repository_vectors
from .github import clone_or_pull_github_repo, fetch_local_git_commits

IGNORED_DIRS = {
    ".git", ".venv", "venv", "node_modules", "__pycache__", "dist",
    "build", ".next", "target", "vendor", ".idea", ".vscode", "coverage", "data"
}

ALLOWED_EXTENSIONS = {
    # Code
    ".py", ".js", ".jsx", ".ts", ".tsx", ".java", ".kt", ".go", ".rs",
    ".c", ".h", ".cpp", ".hpp", ".cs", ".php", ".rb", ".swift", ".sql",
    ".sh", ".ps1", ".bat",
    # Web & Config
    ".html", ".css", ".scss", ".json", ".yaml", ".yml", ".toml", ".xml",
    # Documentation
    ".md", ".txt", ".rst", ".adoc"
}

ALLOWED_FILENAMES = {
    "dockerfile", "makefile", "jenkinsfile", "cmakelists.txt", "procfile", ".env.example"
}

splitter = RecursiveCharacterTextSplitter(
    chunk_size=1200,
    chunk_overlap=180,
    add_start_index=True
)

def build_tree_structure(root_dir: Path, current_dir: Path = None, prefix: str = "", depth: int = 0, max_depth: int = 6) -> tuple[str, dict, dict]:
    """Build ASCII tree, nested dictionary tree, and language/file stats."""
    if current_dir is None:
        current_dir = root_dir

    ascii_lines = []
    nested = {
        "name": current_dir.name if current_dir != root_dir else root_dir.name or "root",
        "type": "directory",
        "path": current_dir.relative_to(root_dir).as_posix() if current_dir != root_dir else ".",
        "children": []
    }
    stats = {"extensions": {}, "total_files": 0, "total_size_bytes": 0}

    if depth >= max_depth:
        return "", nested, stats

    try:
        entries = sorted(list(current_dir.iterdir()), key=lambda e: (not e.is_dir(), e.name.lower()))
    except (PermissionError, OSError):
        return "", nested, stats

    # Filter out ignored directories
    filtered = [e for e in entries if not (e.is_dir() and e.name in IGNORED_DIRS)]
    count = len(filtered)

    for idx, entry in enumerate(filtered):
        is_last = (idx == count - 1)
        connector = "└── " if is_last else "├── "
        child_prefix = "    " if is_last else "│   "

        if entry.is_dir():
            ascii_lines.append(f"{prefix}{connector}{entry.name}/")
            sub_ascii, sub_nested, sub_stats = build_tree_structure(
                root_dir, entry, prefix + child_prefix, depth + 1, max_depth
            )
            nested["children"].append(sub_nested)
            if sub_ascii:
                ascii_lines.append(sub_ascii)
            # Merge stats
            stats["total_files"] += sub_stats["total_files"]
            stats["total_size_bytes"] += sub_stats["total_size_bytes"]
            for ext, ext_count in sub_stats["extensions"].items():
                stats["extensions"][ext] = stats["extensions"].get(ext, 0) + ext_count
        else:
            size = 0
            try:
                size = entry.stat().st_size
            except OSError:
                pass
            ext = entry.suffix.lower() or "no_ext"
            stats["extensions"][ext] = stats["extensions"].get(ext, 0) + 1
            stats["total_files"] += 1
            stats["total_size_bytes"] += size

            ascii_lines.append(f"{prefix}{connector}{entry.name} ({format_size(size)})")
            nested["children"].append({
                "name": entry.name,
                "type": "file",
                "path": entry.relative_to(root_dir).as_posix(),
                "size": size,
                "extension": ext
            })

    return "\n".join(ascii_lines), nested, stats

def format_size(bytes_val: int) -> str:
    """Format bytes into human-readable string."""
    for unit in ['B', 'KB', 'MB', 'GB']:
        if bytes_val < 1024.0:
            return f"{bytes_val:.1f} {unit}" if unit != 'B' else f"{bytes_val} B"
        bytes_val /= 1024.0
    return f"{bytes_val:.1f} TB"

def chunk_file_content(text: str, source: str, repository: str, commit: str, doc_type: str = "code") -> List[Document]:
    """Chunk text content into LangChain Documents with enriched metadata."""
    chunks = splitter.split_text(text)
    return [
        Document(
            page_content=chunk,
            metadata={
                "source": source,
                "repository": repository,
                "commit": commit,
                "chunk": i,
                "type": doc_type,
            }
        )
        for i, chunk in enumerate(chunks)
    ]

def ingest_repository(name: str, path_or_url: str, branch: Optional[str] = None) -> Dict[str, Any]:
    """Ingest a repository from a local filesystem path or clone from GitHub URL."""
    repo_url = None
    is_cloned = False

    # Check if input is a Git URL
    if path_or_url.startswith(("http://", "https://", "git@")):
        clone_res = clone_or_pull_github_repo(repo_url=path_or_url, repo_name=name, branch=branch)
        root = Path(clone_res["path"])
        repo_url = path_or_url
        is_cloned = True
    else:
        root = Path(path_or_url).expanduser().resolve()
        if not root.is_dir():
            raise ValueError(f"Path must be an existing directory: {path_or_url}")

    # Get Git HEAD commit SHA
    try:
        sha = subprocess.check_output(
            ["git", "-C", str(root), "rev-parse", "HEAD"],
            text=True, stderr=subprocess.DEVNULL
        ).strip()
    except Exception:
        sha = "working-tree"

    # Clean existing vectors for this repo to allow clean re-indexing
    delete_repository_vectors(name)

    # Walk directory tree and extract text
    documents = []
    tree_paths = []

    for current, dirs, files in os.walk(root):
        dirs[:] = [d for d in dirs if d not in IGNORED_DIRS]
        for filename in files:
            file_path = Path(current) / filename
            rel_path = file_path.relative_to(root).as_posix()
            tree_paths.append(rel_path)

            ext = file_path.suffix.lower()
            if ext not in ALLOWED_EXTENSIONS and filename.lower() not in ALLOWED_FILENAMES:
                continue

            try:
                if file_path.stat().st_size > MAX_FILE_BYTES:
                    continue
                content = file_path.read_text(encoding="utf-8", errors="replace")
            except OSError:
                continue

            if content.strip():
                documents.extend(chunk_file_content(content, rel_path, name, sha, doc_type="code"))

    # Add chunks to vectorstore
    indexed_count = add_documents(documents)

    # Generate tree view and stats
    ascii_tree, nested_tree, tree_stats = build_tree_structure(root)

    # Upsert repo in database
    upsert_repository(
        name=name,
        path=str(root),
        url=repo_url,
        branch=branch or "main",
        last_commit_sha=sha,
        total_files=len(tree_paths)
    )

    # Extract recent commit history and populate logbook
    commits = fetch_local_git_commits(str(root), limit=30)
    for c in commits:
        record_commit(
            repository=name,
            sha=c["sha"],
            author=c["author"],
            email=c["email"],
            message=c["message"],
            timestamp=c["timestamp"],
            files_added=c["files_added"],
            files_modified=c["files_modified"],
            files_removed=c["files_removed"],
        )

    # Record initial logbook entry
    if commits:
        latest = commits[0]
        record_change_log(
            repository=name,
            sha=latest["sha"],
            author=latest["author"],
            summary=f"Indexed repository '{name}' at commit {latest['sha'][:8]}: {latest['message']}",
            diff_stat=f"+{len(latest['files_added'])} ~{len(latest['files_modified'])} -{len(latest['files_removed'])}"
        )

    return {
        "repository": name,
        "path": str(root),
        "url": repo_url,
        "is_cloned": is_cloned,
        "commit": sha,
        "total_files": len(tree_paths),
        "indexed_chunks": indexed_count,
        "ascii_tree": ascii_tree[:15000],  # truncated if huge
        "nested_tree": nested_tree,
        "tree_stats": tree_stats,
        "commits_discovered": len(commits),
    }

def ingest_document(file_path: str, repository: str, original_filename: Optional[str] = None) -> Dict[str, Any]:
    """Parse and index documentation files (.pdf, .docx, .md, .txt)."""
    p = Path(file_path).expanduser().resolve()
    if not p.is_file():
        raise ValueError(f"Document file not found: {file_path}")

    filename = original_filename or p.name
    size = p.stat().st_size
    if size > MAX_FILE_BYTES * 15:
        raise ValueError("Document exceeds maximum size limit (7.5 MB).")

    ext = p.suffix.lower()
    content = ""

    if ext in {".md", ".txt", ".rst"}:
        content = p.read_text(encoding="utf-8", errors="replace")
    elif ext == ".pdf":
        reader = PdfReader(str(p))
        pages_text = []
        for i, page in enumerate(reader.pages):
            txt = page.extract_text() or ""
            if txt.strip():
                pages_text.append(f"--- Page {i+1} ---\n{txt}")
        content = "\n\n".join(pages_text)
    elif ext == ".docx":
        doc = WordDocument(str(p))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        content = "\n\n".join(paragraphs)
    else:
        raise ValueError(f"Unsupported document format: {ext}. Supported formats: .pdf, .docx, .md, .txt")

    if not content.strip():
        raise ValueError("Extracted document content is empty.")

    # Chunk and store
    chunks = chunk_file_content(content, filename, repository, "document", doc_type="document")
    add_documents(chunks)

    # Record in database
    record_document(
        repository=repository,
        filename=filename,
        file_path=str(p),
        file_size=size,
        chunk_count=len(chunks),
        file_type=ext.removeprefix(".")
    )

    return {
        "repository": repository,
        "filename": filename,
        "file_size": size,
        "indexed_chunks": len(chunks),
        "file_type": ext.removeprefix("."),
    }
